"""Research toolkit: figures, tables, paper pipeline and experiment runs."""
import base64
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from docx import Document  # noqa: E402

from life.tools.research import (  # noqa: E402
    bar, build_paper, line, markdown_to_html, run_experiment,
    table_to_html, table_to_latex)


class FiguresTest(unittest.TestCase):
    def test_line_and_bar_write_png(self):
        with tempfile.TemporaryDirectory() as tmp:
            png = line([0, 1, 2], {"训练": [0.2, 0.5, 0.8], "验证": [0.1, 0.4, 0.7]},
                       title="学习曲线", path=Path(tmp) / "curve.png")
            self.assertTrue(png.is_file() and png.stat().st_size > 0)
            self.assertEqual(png.suffix, ".png")
            bars = bar(["基线", "本方法"], [0.61, 0.88], title="对比",
                       path=Path(tmp) / "bar.png", horizontal=True)
            self.assertTrue(bars.is_file() and bars.stat().st_size > 0)


class TablesTest(unittest.TestCase):
    def test_three_line_html_and_latex(self):
        html = table_to_html(["方法", "准确率"], [["基线", "0.61"], ["本方法", "0.88"]],
                             caption="表1 结果")
        self.assertIn('class="three-line"', html)
        self.assertIn("表1 结果", html)
        self.assertIn("本方法", html)
        latex = table_to_latex(["方法", "准确率"], [["基线", "0.61"]], label="tab:r")
        self.assertIn("\\toprule", latex)
        self.assertIn("\\bottomrule", latex)
        self.assertIn("\\label{tab:r}", latex)


class PaperTest(unittest.TestCase):
    MARKDOWN = (
        "# 标题\n\n"
        "这是**加粗**段落，含 `代码` 与公式。\n\n"
        "## 方法\n\n"
        "- 步骤一\n- 步骤二\n\n"
        "| 方法 | 准确率 |\n| --- | --- |\n| 基线 | 0.61 |\n| 本方法 | 0.88 |\n\n"
        "```python\nprint('hello')\n```\n\n"
        "![图1 结构](images/fig1.png)\n"
    )

    def test_markdown_to_html_blocks(self):
        html = markdown_to_html(self.MARKDOWN, title="论文")
        for token in ("<h1>标题</h1>", "<strong>加粗</strong>", "<code>代码</code>",
                      "<ul>", "three-line", "<pre><code", "<figure>", "<title>论文</title>"):
            self.assertIn(token, html)

    def test_build_paper_writes_html_and_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = build_paper(self.MARKDOWN, tmp, title="论文", stem="draft")
            self.assertTrue(paths["html"].is_file())
            self.assertTrue(paths["docx"].is_file())
            self.assertGreater(paths["docx"].stat().st_size, 0)
            self.assertTrue(paths["docx"].read_bytes()[:2] == b"PK")


class PaperLayoutTest(unittest.TestCase):
    """Layout regressions: figure rows, indents, heading colours, links, footer."""

    PNG_1PX = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAgAAAAICAAAAADhZOFXAAAAI0lEQVR4nGNggAERKGDQgAIGGyh"
        "gCIAChhQoYKiAAoYeKAAAihcRgUPxg80AAAAASUVORK5CYII=")

    def test_consecutive_figures_form_one_row(self):
        html = markdown_to_html("![甲](a.png)\n![乙](b.png){width=30%}\n")
        self.assertIn('class="fig-group"', html)
        self.assertIn("图 1(a)", html)
        self.assertIn("图 1(b)", html)
        # 30% explicit + equal share of the rest fills the row.
        self.assertIn('data-frac="0.3000"', html)
        self.assertIn('data-frac="0.7000"', html)
        html = markdown_to_html("![甲](a.png)\n![乙](b.png)\n![丙](c.png)\n")
        self.assertEqual(html.count('data-frac="0.3333"'), 3)
        # A blank line breaks the group: separate numbered figures again.
        html = markdown_to_html("![甲](a.png)\n\n![乙](b.png)\n")
        self.assertNotIn('class="fig-group"', html)
        self.assertIn("图 1", html)
        self.assertIn("图 2", html)

    def test_docx_figure_group_shares_one_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in ("a.png", "b.png"):
                (Path(tmp) / name).write_bytes(self.PNG_1PX)
            paths = build_paper("![甲](a.png)\n![乙](b.png)\n\n正文一段。\n", tmp, title="论文")
            document = Document(paths["docx"])
            self.assertEqual(len(document.tables), 1)
            table = document.tables[0]
            self.assertEqual(len(table.rows), 1)
            self.assertEqual(len(table.columns), 2)
            cells = table.rows[0].cells
            for cell in cells:
                self.assertIn("<w:drawing>", cell._tc.xml)
            widths = sorted(cell.width.cm for cell in cells)
            self.assertAlmostEqual(widths[0], 8.1, delta=0.3)
            self.assertIn('w:val="none"', table._tbl.tblPr.xml)
            self.assertIn("图 1(a)", cells[0].text)
            self.assertIn("图 1(b)", cells[1].text)

    def test_single_figure_stays_a_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "a.png").write_bytes(self.PNG_1PX)
            paths = build_paper("![甲](a.png)\n", tmp, title="论文")
            document = Document(paths["docx"])
            self.assertEqual(len(document.tables), 0)
            self.assertTrue(document.inline_shapes)

    def test_docx_styles_indent_and_page_number(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = build_paper("# 标题\n\n正文段落，检查首行缩进。\n", tmp, title="论文")
            document = Document(paths["docx"])
            self.assertEqual(str(document.styles["Heading 1"].font.color.rgb), "1B2430")
            self.assertEqual(document.styles["Title"].font.size.pt, 17.0)
            body = next(p for p in document.paragraphs if p.text == "正文段落，检查首行缩进。")
            self.assertAlmostEqual(body.paragraph_format.first_line_indent.pt, 21.0)
            footer = document.sections[0].footer
            self.assertIn("PAGE", footer.paragraphs[0]._p.xml)

    def test_docx_hyperlinks_are_real_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown = "参见[官方文档](https://example.com/docs)与[本地](javascript:alert(1))。\n"
            paths = build_paper(markdown, tmp, title="论文")
            document = Document(paths["docx"])
            links = document.element.body.xpath(".//w:hyperlink")
            self.assertEqual(len(links), 1)
            self.assertEqual(links[0].xpath("./w:r/w:t/text()"), ["官方文档"])
            targets = [rel.target_ref for rel in document.part.rels.values()]
            self.assertTrue(any("example.com" in target for target in targets))
            self.assertFalse(any("javascript" in target for target in targets))


class ExperimentTest(unittest.TestCase):
    def test_run_records_stdout_and_result(self):
        def work():
            print("epoch 1 loss 0.5")
            return {"accuracy": 0.88}

        with tempfile.TemporaryDirectory() as tmp:
            run = run_experiment("smoke", work, out_dir=tmp, params={"seed": 7})
            self.assertEqual(run.result, {"accuracy": 0.88})
            self.assertIn("epoch 1", run.stdout)
            self.assertTrue((run.run_dir / "stdout.txt").is_file())
            meta = json.loads((run.run_dir / "meta.json").read_text(encoding="utf-8"))
            self.assertEqual(meta["params"], {"seed": 7})
            self.assertEqual(meta["error"], "")

    def test_run_captures_failure(self):
        def boom():
            raise ValueError("nope")

        with tempfile.TemporaryDirectory() as tmp:
            run = run_experiment("fail", boom, out_dir=tmp)
            self.assertIn("ValueError", run.error)
            self.assertIsNone(run.result)


class CliTest(unittest.TestCase):
    def _write_experiment(self, tmp: str):
        script = Path(tmp) / "exp.py"
        script.write_text(
            "def main(seed=0):\n"
            "    print('running seed', seed)\n"
            "    return {'seed': seed, 'accuracy': 0.9}\n",
            encoding="utf-8",
        )
        spec = Path(tmp) / "spec.json"
        spec.write_text(
            json.dumps({"script": str(script), "name": "exp", "params": {"seed": 3}}),
            encoding="utf-8",
        )
        return spec

    def test_experiment_command_runs_script_when_enabled(self):
        from life.tools.research.__main__ import main

        with tempfile.TemporaryDirectory() as tmp:
            spec = self._write_experiment(tmp)
            env = {"LIFE_RESEARCH_ALLOW_EXEC": "1", "LIFE_RESEARCH_DIR": tmp}
            with mock.patch.dict(os.environ, env):
                self.assertEqual(main(["experiment", str(spec), "-o", tmp]), 0)
            runs = sorted(p for p in Path(tmp).glob("*exp*") if p.is_dir())
            self.assertTrue(runs)
            result = json.loads((runs[0] / "result.json").read_text(encoding="utf-8"))
            self.assertEqual(result, {"seed": 3, "accuracy": 0.9})

    def test_experiment_command_refuses_without_opt_in(self):
        """Executing a model-authored script must not be on by default."""
        from life.tools.research.__main__ import main

        with tempfile.TemporaryDirectory() as tmp:
            spec = self._write_experiment(tmp)
            environ = dict(os.environ)
            environ.pop("LIFE_RESEARCH_ALLOW_EXEC", None)
            with mock.patch.dict(os.environ, environ, clear=True):
                with self.assertRaises(SystemExit):
                    main(["experiment", str(spec), "-o", tmp])

    def test_experiment_command_refuses_scripts_outside_allowed_dir(self):
        from life.tools.research.__main__ import main

        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as elsewhere:
            spec = self._write_experiment(tmp)
            env = {"LIFE_RESEARCH_ALLOW_EXEC": "1", "LIFE_RESEARCH_DIR": elsewhere}
            with mock.patch.dict(os.environ, env):
                with self.assertRaises(SystemExit):
                    main(["experiment", str(spec), "-o", tmp])

    def test_paper_command_writes_outputs(self):
        from life.tools.research.__main__ import main

        with tempfile.TemporaryDirectory() as tmp:
            md = Path(tmp) / "paper.md"
            md.write_text("# 标题\n\n正文。\n", encoding="utf-8")
            self.assertEqual(main(["paper", str(md), "-o", tmp, "--title", "论文"]), 0)
            self.assertTrue((Path(tmp) / "paper.html").is_file())
            self.assertTrue((Path(tmp) / "paper.docx").is_file())


if __name__ == "__main__":
    unittest.main()
