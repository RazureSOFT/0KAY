"""Markdown -> HTML -> docx publishing pipeline.

Renders a paper from Markdown into (a) a styled HTML review document and (b) a
Word document with a real title block, sectioned headings, justified body text,
centred figures with numbered captions and three-line tables.

Design notes
------------
* The HTML is parsed with the tolerant stdlib :class:`html.parser.HTMLParser`
  rather than an XML parser.  Model-authored Markdown routinely carries a stray
  ``&`` or ``<``; an XML parse error used to collapse the whole document into a
  single plain-text paragraph (losing every figure and all layout).
* Layout is author-controllable from the Markdown alone:
    - a leading ``---`` YAML-ish front-matter block (``title``, ``authors``,
      ``affiliation``, ``abstract``, ``keywords``);
    - figures ``![caption](path){width=60%}`` (also px/in/cm);
    - consecutive figure lines with no blank line between them form one row
      (``![a](1.png)`` + ``![b](2.png)`` → 图1(a)/(b) side by side; widths are
      honoured, missing ones share the remaining row equally);
    - three-line (booktabs) tables.
  Relative image paths resolve against the Markdown file's directory, the
  output directory and the current directory, in that order.
"""
from __future__ import annotations

import re
from html import escape
from html.parser import HTMLParser
from pathlib import Path

# --- front matter -------------------------------------------------------------

def parse_front_matter(text: str) -> tuple[dict[str, str], str]:
    """Split a leading ``---`` metadata block from the Markdown body."""
    raw = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = raw.split("\n")
    start = 0
    while start < len(lines) and not lines[start].strip():
        start += 1
    if start >= len(lines) or lines[start].strip() != "---":
        return {}, raw
    end = None
    for index in range(start + 1, len(lines)):
        if lines[index].strip() in ("---", "..."):
            end = index
            break
    if end is None:
        return {}, raw
    meta: dict[str, str] = {}
    for line in lines[start + 1:end]:
        if not line.strip() or line.lstrip().startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        meta[key.strip().lower()] = value.strip().strip('"').strip("'")
    return meta, "\n".join(lines[end + 1:])


# --- tolerant HTML tree -------------------------------------------------------

_VOID = {"br", "hr", "img", "meta", "link", "input", "col", "area", "base",
         "source", "track", "wbr"}


class _Node:
    __slots__ = ("tag", "attrs", "children", "text")

    def __init__(self, tag: str, attrs=None, text: str = ""):
        self.tag = tag
        self.attrs = {key: (value or "") for key, value in (attrs or [])}
        self.children: list["_Node"] = []
        self.text = text


class _TreeBuilder(HTMLParser):
    """Build a simple element tree; tolerant of unescaped ``&``/``<`` and HTML entities."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = _Node("#root")
        self._stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = _Node(tag, attrs)
        self._stack[-1].children.append(node)
        if tag not in _VOID:
            self._stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self._stack[-1].children.append(_Node(tag, attrs))

    def handle_endtag(self, tag):
        for index in range(len(self._stack) - 1, 0, -1):
            if self._stack[index].tag == tag:
                del self._stack[index:]
                return

    def handle_data(self, data):
        if data:
            self._stack[-1].children.append(_Node("#text", text=data))


def _parse_html(html: str) -> _Node:
    builder = _TreeBuilder()
    builder.feed(html or "")
    builder.close()
    return builder.root


def _walk(node: _Node, tag: str):
    for child in node.children:
        if child.tag == tag:
            yield child
        yield from _walk(child, tag)


def _first(node: _Node, tag: str) -> _Node | None:
    return next(_walk(node, tag), None)


# --- markdown -> html ---------------------------------------------------------

_IMG_INLINE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)(?:\{([^}]*)\})?")
_LINK_INLINE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _is_block_start(line: str) -> bool:
    stripped = line.strip()
    return bool(
        re.match(r"^(#{1,6})\s+", stripped)
        or stripped.startswith("```")
        or stripped.startswith(">")
        or re.match(r"^([-*+]|\d+\.)\s+", stripped)
        or re.match(r"^!\[[^\]]*\]\([^)]+\)(?:\{[^}]*\})?\s*$", stripped)
        or re.match(r"^\s*([-*_])\1{2,}\s*$", stripped)
    )


def _image_attrs(attrs: str) -> str:
    match = re.search(r"width\s*=\s*([0-9.]+)\s*(%|px|in|cm)?", attrs or "")
    return f' data-width="{match.group(1)}{match.group(2) or ""}"' if match else ""


# --- figure groups ------------------------------------------------------------

_TEXT_WIDTH_CM = 16.2  # A4 width minus the 2.4cm side margins


def _width_fraction(attrs: str) -> float | None:
    """Interpret a figure ``{width=...}`` spec as a fraction of the text width."""
    match = re.search(r"width\s*=\s*([0-9.]+)\s*(%|px|in|cm)?", attrs or "")
    if not match:
        return None
    value = float(match.group(1))
    unit = match.group(2) or "in"
    if unit == "%":
        return value / 100.0
    if unit == "px":
        return value / 96.0 * 2.54 / _TEXT_WIDTH_CM
    if unit == "cm":
        return value / _TEXT_WIDTH_CM
    return value * 2.54 / _TEXT_WIDTH_CM


def _group_fractions(attrs_list: list[str]) -> list[float]:
    """Split the text width among the figures of one row.

    Explicit widths are honoured; figures without one share the remaining width
    equally, and an overflow is scaled back so the row always fits the page.
    """
    count = max(1, len(attrs_list))
    fracs = [_width_fraction(attrs) for attrs in attrs_list]
    if all(frac is None for frac in fracs):
        return [1.0 / count] * len(fracs)
    given = sum(frac for frac in fracs if frac is not None)
    missing = [index for index, frac in enumerate(fracs) if frac is None]
    if missing and given < 1.0:
        share = (1.0 - given) / len(missing)
        fracs = [share if frac is None else frac for frac in fracs]
        given = 1.0
    elif missing:
        fracs = [1.0 / count if frac is None else frac for frac in fracs]
        given = sum(fracs)
    if given > 1.0:
        fracs = [frac / given for frac in fracs]
    return fracs


def _img_tag(alt: str, src: str, attrs: str = "") -> str:
    return (f'<img src="{escape(src, quote=True)}" alt="{escape(alt, quote=True)}"'
            f'{_image_attrs(attrs)}/>')


def _inline(text: str) -> str:
    text = escape(text, quote=False)
    text = _IMG_INLINE.sub(
        lambda m: _img_tag(m.group(1), m.group(2), m.group(3) or ""), text)
    text = _LINK_INLINE.sub(
        lambda m: f'<a href="{escape(m.group(2), quote=True)}">{m.group(1)}</a>', text)
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", text)
    return text


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _table_html(header: list[str], rows: list[list[str]]) -> str:
    from .tables import table_to_html

    return table_to_html(header, rows)


def _title_block(meta: dict[str, str]) -> str:
    title = str(meta.get("title") or "").strip()
    authors = str(meta.get("authors") or meta.get("author") or "").strip()
    affiliation = str(meta.get("affiliation") or "").strip()
    abstract = str(meta.get("abstract") or "").strip()
    keywords = str(meta.get("keywords") or "").strip()
    if not any((title, authors, abstract, keywords)):
        return ""
    parts = ['<header class="doc-header">']
    if title:
        parts.append(f'<h1 class="doc-title">{escape(title)}</h1>')
    if authors:
        line = escape(authors)
        if affiliation:
            line += f'<br/><span class="doc-affil">{escape(affiliation)}</span>'
        parts.append(f'<p class="doc-authors">{line}</p>')
    if abstract:
        parts.append(f'<div class="doc-abstract"><p><strong>摘要　</strong>{_inline(abstract)}</p></div>')
    if keywords:
        parts.append(f'<p class="doc-keywords"><strong>关键词：</strong>{_inline(keywords)}</p>')
    parts.append("</header>")
    return "".join(parts)


def markdown_to_html(markdown_text: str, *, title: str = "", css: str = "",
                     meta: dict[str, str] | None = None) -> str:
    """Convert a Markdown document into a styled, well-formed HTML document."""
    if meta is None:
        meta, markdown_text = parse_front_matter(markdown_text)
    if title and not meta.get("title"):
        meta["title"] = title

    lines = str(markdown_text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    parts: list[str] = []
    header = _title_block(meta)
    if header:
        parts.append(header)

    # Drop a leading `# <title>` that merely repeats the front-matter title.
    doc_title = str(meta.get("title") or "").strip()
    for position, line in enumerate(lines):
        if not line.strip():
            continue
        if doc_title and re.match(r"^#\s+" + re.escape(doc_title) + r"\s*$", line.strip()):
            lines[position] = ""
        break

    figure_no = 0
    table_no = 0
    index, total = 0, len(lines)
    while index < total:
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            index += 1
            buffer: list[str] = []
            while index < total and not lines[index].strip().startswith("```"):
                buffer.append(lines[index])
                index += 1
            index += 1
            code = escape("\n".join(buffer))
            cls = f' class="language-{escape(language)}"' if language else ""
            parts.append(f"<pre><code{cls}>{code}</code></pre>")
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            parts.append(f"<h{level}>{_inline(heading.group(2).strip())}</h{level}>")
            index += 1
            continue

        if re.match(r"^([-*_])\1{2,}$", stripped):
            parts.append("<hr/>")
            index += 1
            continue

        if (index + 1 < total and "|" in line
                and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1] or "")):
            header_row = _split_row(line)
            index += 2
            rows = []
            while index < total and "|" in lines[index] and lines[index].strip():
                rows.append(_split_row(lines[index]))
                index += 1
            table_no += 1
            parts.append(_table_html(header_row, rows))
            continue

        if re.match(r"^\s*([-*+]|\d+\.)\s+", line):
            ordered = bool(re.match(r"^\s*\d+\.", line))
            items = []
            while index < total and re.match(r"^\s*([-*+]|\d+\.)\s+", lines[index]):
                items.append(re.sub(r"^\s*([-*+]|\d+\.)\s+", "", lines[index]).strip())
                index += 1
            tag = "ol" if ordered else "ul"
            body = "".join(f"<li>{_inline(item)}</li>" for item in items)
            parts.append(f"<{tag}>{body}</{tag}>")
            continue

        if line.startswith(">"):
            quoted = []
            while index < total and lines[index].startswith(">"):
                quoted.append(lines[index][1:].strip())
                index += 1
            parts.append(f"<blockquote><p>{_inline(' '.join(quoted))}</p></blockquote>")
            continue

        image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)(?:\{([^}]*)\})?\s*$", stripped)
        if image:
            # Consecutive figure lines (no blank line between) share one row.
            group: list[re.Match] = []
            while index < total:
                member = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)(?:\{([^}]*)\})?\s*$",
                                  lines[index].strip())
                if not member:
                    break
                group.append(member)
                index += 1
            figure_no += 1
            if len(group) == 1:
                alt, src, extra = group[0].group(1), group[0].group(2), group[0].group(3) or ""
                caption = _inline(alt) if alt else ""
                cap = (f'<figcaption><strong>图 {figure_no}</strong>'
                       f'{("　" + caption) if caption else ""}</figcaption>')
                parts.append(f'<figure>{_img_tag(alt, src, extra)}{cap}</figure>')
                continue
            fracs = _group_fractions([member.group(3) or "" for member in group])
            items = []
            for position, member in enumerate(group):
                alt, src, extra = member.group(1), member.group(2), member.group(3) or ""
                letter = chr(ord("a") + position)
                caption = _inline(alt) if alt else ""
                cap = (f'<figcaption><strong>图 {figure_no}({letter})</strong>'
                       f'{("　" + caption) if caption else ""}</figcaption>')
                percent = fracs[position] * 100
                items.append(f'<span class="fig-item" data-frac="{fracs[position]:.4f}"'
                             f' style="flex:0 1 {percent:.1f}%">'
                             f'{_img_tag(alt, src, extra)}{cap}</span>')
            parts.append(f'<figure class="fig-group">{"".join(items)}</figure>')
            continue

        if not stripped:
            index += 1
            continue

        buffer = [stripped]
        index += 1
        while index < total and lines[index].strip() and not _is_block_start(lines[index]):
            buffer.append(lines[index].strip())
            index += 1
        parts.append(f"<p>{_inline(' '.join(buffer))}</p>")

    return _document("\n".join(parts), str(meta.get("title") or title), css)


def _document(body: str, title: str, css: str) -> str:
    head = ['<meta charset="utf-8"/>', '<meta name="viewport" content="width=device-width, initial-scale=1"/>']
    if title:
        head.append(f"<title>{escape(title)}</title>")
    style = css if css != "" else _ACADEMIC_CSS
    if style:
        head.append(f"<style>{style}</style>")
    return f"<!DOCTYPE html><html><head>{''.join(head)}</head><body>{body}</body></html>"


# --- html -> docx -------------------------------------------------------------

_ACADEMIC_CSS = """
:root { --ink:#1b2430; --muted:#5b6b7a; --rule:#c9d2dc; --accent:#3a6ea5; }
@page { size: A4; margin: 2.2cm 2.4cm; }
body { max-width: 820px; margin: 0 auto; padding: 40px 22px; color: var(--ink);
  font-family: "Times New Roman", "Songti SC", "SimSun", "Noto Serif CJK SC", serif;
  font-size: 16px; line-height: 1.75; }
header.doc-header { text-align: center; margin: 0 0 30px; }
h1.doc-title { font-size: 1.9rem; line-height: 1.35; margin: 0 0 14px; font-weight: 700; }
p.doc-authors { margin: 0 0 6px; font-size: 1.02rem; }
.doc-affil { color: var(--muted); font-size: .95rem; }
.doc-abstract { text-align: justify; margin: 18px auto; padding: 14px 18px;
  background: #f6f8fb; border: 1px solid var(--rule); border-radius: 10px; font-size: .96rem; }
.doc-abstract p { margin: 0; }
p.doc-keywords { margin: 4px auto 20px; font-size: .95rem; }
h1, h2, h3, h4 { line-height: 1.4; }
h1 { font-size: 1.5rem; margin: 1.6em 0 .6em; }
h2 { font-size: 1.28rem; margin: 1.5em 0 .5em; border-bottom: 1px solid var(--rule); padding-bottom: 5px; }
h3 { font-size: 1.09rem; margin: 1.2em 0 .4em; }
p { text-align: justify; margin: .7em 0; }
table.three-line { border-collapse: collapse; margin: 1.2em auto; font-size: .96rem; }
table.three-line caption { caption-side: top; font-weight: 700; margin-bottom: 8px; }
table.three-line th, table.three-line td { padding: 6px 16px; text-align: center; }
table.three-line thead tr { border-top: 2px solid var(--ink); border-bottom: 1px solid var(--ink); }
table.three-line tbody tr:last-child { border-bottom: 2px solid var(--ink); }
pre { background: #f5f7fa; border: 1px solid var(--rule); border-radius: 8px;
  padding: 12px 14px; overflow-x: auto; font-size: .9rem; line-height: 1.5; }
code { font-family: "Cascadia Code", Consolas, monospace; font-size: .92em; }
figure { text-align: center; margin: 1.6em 0; }
figure img { max-width: 100%; height: auto; }
figcaption { color: var(--muted); font-size: .9rem; margin-top: 8px; }
figure.fig-group { display: flex; gap: 16px; justify-content: center; align-items: flex-start; }
figure.fig-group .fig-item { display: block; min-width: 0; }
figure.fig-group .fig-item img { width: 100%; height: auto; }
figure.fig-group figcaption { margin-top: 6px; }
blockquote { border-left: 3px solid var(--accent); margin: 1em 0; padding: 4px 16px;
  color: var(--muted); }
hr { border: none; border-top: 1px solid var(--rule); margin: 1.6em 0; }
"""


def _extract_body(html: str) -> str:
    match = re.search(r"<body[^>]*>(.*)</body>", html or "", flags=re.S | re.I)
    return match.group(1) if match else (html or "")


def _set_east_asian(style_or_run, family: str) -> None:
    from docx.oxml.ns import qn

    rpr = style_or_run.font._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), family)


def _strip_style_border(style) -> None:
    """Drop a decorative paragraph border the template's Title style may carry."""
    from docx.oxml.ns import qn

    p_pr = style.element.find(qn("w:pPr"))
    if p_pr is not None:
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)


def _add_page_number(section) -> None:
    """Centred ``PAGE`` field in the footer."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer = section.footer
    if footer.is_linked_to_previous:
        footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), r" PAGE \* MERGEFORMAT ")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")  # 9pt
    r_pr.append(size)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _style_document(document) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    ink = RGBColor(0x1B, 0x24, 0x30)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = ink
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_east_asian(normal, "宋体")

    # The bundled template's headings are Word-blue; papers need ink-black ones.
    for name, size in (("Title", 17), ("Heading 1", 15), ("Heading 2", 13),
                       ("Heading 3", 11.5), ("Heading 4", 10.5)):
        try:
            style = document.styles[name]
        except KeyError:
            continue
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.name = "Times New Roman"
        style.font.color.rgb = ink
        _set_east_asian(style, "黑体")
        if name == "Title":
            style.paragraph_format.space_after = Pt(12)
            _strip_style_border(style)

    _add_page_number(section)


def _add_runs(paragraph, node: _Node, *, bold=False, italic=False, code=False) -> None:
    from docx.shared import Pt

    for child in node.children:
        tag = child.tag
        if tag == "#text":
            if child.text:
                run = paragraph.add_run(child.text)
                if bold:
                    run.bold = True
                if italic:
                    run.italic = True
                if code:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9.5)
        elif tag == "br":
            paragraph.add_run().add_break()
        elif tag in ("strong", "b"):
            _add_runs(paragraph, child, bold=True, italic=italic, code=code)
        elif tag in ("em", "i"):
            _add_runs(paragraph, child, bold=bold, italic=True, code=code)
        elif tag == "code":
            _add_runs(paragraph, child, bold=bold, italic=italic, code=True)
        elif tag == "a":
            _add_link(paragraph, child, bold=bold, italic=italic)
        elif tag == "img":
            _add_inline_image(paragraph, child)
        else:
            _add_runs(paragraph, child, bold=bold, italic=italic, code=code)


def _node_text(node: _Node) -> str:
    return "".join(child.text for child in _walk(node, "#text"))


def _add_link(paragraph, node: _Node, *, bold=False, italic=False) -> None:
    """Render ``<a>`` as a real hyperlink; unknown schemes degrade to plain text."""
    href = str(node.attrs.get("href") or "").strip()
    text = _node_text(node)
    if not text:
        return
    if not re.match(r"^(https?://|mailto:)", href, re.I):
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        return

    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relation = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3a6ea5")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    if italic:
        r_pr.append(OxmlElement("w:i"))
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline_image(paragraph, node: _Node) -> None:
    """Inline ``<img>`` inside a paragraph; block figures go through _emit_image."""
    src = str(node.attrs.get("src") or "")
    run = paragraph.add_run()
    resolved = _resolve_image(src, _INLINE_BASE_DIRS)
    if resolved is None:
        run.text = f"[缺失图片：{src}]"
        return
    try:
        run.add_picture(str(resolved), width=_image_width(node.attrs, default_inches=2.5))
    except Exception:
        run.text = f"[图片无法插入：{src}]"


# Base dirs for resolving inline images; html_to_docx pins these for one render pass.
_INLINE_BASE_DIRS: tuple = ()


def _resolve_image(src: str, base_dirs) -> Path | None:
    candidate = Path(str(src).strip())
    if candidate.is_absolute():
        return candidate if candidate.is_file() else None
    for base in base_dirs:
        resolved = Path(base) / candidate
        if resolved.is_file():
            return resolved
    return candidate if candidate.is_file() else None


def _image_width(attrs: dict, *, default_inches: float = 5.9):
    from docx.shared import Cm, Inches

    raw = str(attrs.get("data-width") or "").strip()
    match = re.match(r"([0-9.]+)\s*(%|px|in|cm)?$", raw)
    if not match:
        return Inches(default_inches)
    value = float(match.group(1))
    unit = match.group(2) or "in"
    if unit == "%":
        return Inches(6.3 * value / 100.0)
    if unit == "px":
        return Inches(value / 96.0)
    if unit == "cm":
        return Cm(value)
    return Inches(value)


def _emit_image(document, node: _Node, base_dirs) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    src = str(node.attrs.get("src") or "")
    alt = str(node.attrs.get("alt") or "")
    resolved = _resolve_image(src, base_dirs)
    if resolved is None:
        paragraph = document.add_paragraph(f"[缺失图片：{src or alt}]")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    try:
        document.add_picture(str(resolved), width=_image_width(node.attrs))
        paragraph = document.paragraphs[-1]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
    except Exception:
        paragraph = document.add_paragraph(f"[图片无法插入：{src}]")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _set_first_line_indent(paragraph) -> None:
    """Two-character first-line indent (CJK convention; font-relative in Word)."""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    paragraph.paragraph_format.first_line_indent = Pt(21)
    ind = paragraph._p.get_or_add_pPr().find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), "200")


def _emit_figure_group(document, node: _Node, base_dirs) -> None:
    """Render a fig-group as a borderless table so the images share one row."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    items = [child for child in node.children
             if child.tag == "span" and "fig-item" in str(child.attrs.get("class") or "")]
    if not items:
        return
    count = len(items)
    fractions = []
    for item in items:
        try:
            fraction = float(str(item.attrs.get("data-frac") or ""))
        except ValueError:
            fraction = 0.0
        fractions.append(fraction if fraction > 0 else 1.0 / count)
    total = sum(fractions)
    if total <= 0:
        fractions, total = [1.0 / count] * count, 1.0
    if total > 1.0:
        fractions = [fraction / total for fraction in fractions]

    table = document.add_table(rows=1, cols=count)
    table.alignment = 1  # centered
    table.autofit = False
    for position, item in enumerate(items):
        image = _first(item, "img")
        caption = _first(item, "figcaption")
        cell = table.cell(0, position)
        cell.width = Cm(fractions[position] * _TEXT_WIDTH_CM)
        paragraph = cell.paragraphs[0]
        if image is not None and image.attrs.get("src"):
            src = str(image.attrs.get("src") or "")
            resolved = _resolve_image(src, base_dirs)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            if resolved is None:
                paragraph.add_run(f"[缺失图片：{src}]")
            else:
                try:
                    # The default cell margins act as the gutter between images.
                    paragraph.add_run().add_picture(
                        str(resolved), width=Cm(fractions[position] * _TEXT_WIDTH_CM - 0.4))
                except Exception:
                    paragraph.add_run(f"[图片无法插入：{src}]")
        if caption is not None:
            paragraph = cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(paragraph, caption)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
    _no_borders(table)
    # A spacer keeps the group table from touching a following table.
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer.add_run("").font.size = Pt(4)


def _emit_caption(document, node: _Node) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    _add_runs(paragraph, node)
    for run in paragraph.runs:
        run.font.size = Pt(9.5)


def _cell_border(cell, edge: str, size: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:color"), "1b2430")
    element.set(qn("w:space"), "0")


def _no_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none")
        element.set(qn("w:sz"), "0")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _three_line_table(table) -> None:
    _no_borders(table)
    for cell in table.rows[0].cells:
        _cell_border(cell, "top", 12)
        _cell_border(cell, "bottom", 6)
    for cell in table.rows[-1].cells:
        _cell_border(cell, "bottom", 12)


def _emit_table(document, node: _Node) -> None:
    rows = list(_walk(node, "tr"))
    if not rows:
        return
    caption = _first(node, "caption")
    if caption is not None:
        _emit_caption(document, caption)
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    cells = [cell for cell in rows[0].children if cell.tag in ("th", "td")]
    table = document.add_table(rows=len(rows), cols=max(1, len(cells)))
    table.alignment = 1  # center
    for row_index, row in enumerate(rows):
        cells = [cell for cell in row.children if cell.tag in ("th", "td")]
        for column, cell in enumerate(cells):
            paragraph = table.cell(row_index, column).paragraphs[0]
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(paragraph, cell)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    try:
        _three_line_table(table)
    except Exception:
        table.style = "Table Grid"


def _emit_block(document, node: _Node, *, base_dirs=()) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    tag = node.tag
    classes = str(node.attrs.get("class") or "")

    if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
        if "doc-title" in classes:
            paragraph = document.add_heading("", level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(paragraph, node)
            return
        paragraph = document.add_heading("", level=min(int(tag[1]), 4))
        _add_runs(paragraph, node)
        return

    if tag == "p":
        paragraph = document.add_paragraph()
        if "doc-authors" in classes:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif "doc-keywords" in classes:
            paragraph.paragraph_format.space_before = Pt(6)
        else:
            _set_first_line_indent(paragraph)
        _add_runs(paragraph, node)
        return

    if tag == "header" or tag == "section":
        for child in node.children:
            _emit_block(document, child, base_dirs=base_dirs)
        return

    if tag == "div":
        if "doc-abstract" in classes:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.7)
            for child in node.children:
                if child.tag == "p":
                    _add_runs(paragraph, child)
                elif child.tag == "#text":
                    paragraph.add_run(child.text)
            return
        for child in node.children:
            _emit_block(document, child, base_dirs=base_dirs)
        return

    if tag == "figure":
        if "fig-group" in classes:
            _emit_figure_group(document, node, base_dirs)
            return
        image = _first(node, "img")
        if image is not None and image.attrs.get("src"):
            _emit_image(document, image, base_dirs)
        caption = _first(node, "figcaption")
        if caption is not None:
            _emit_caption(document, caption)
        return

    if tag == "img":
        _emit_image(document, node, base_dirs)
        return

    if tag in ("ul", "ol"):
        style = "List Number" if tag == "ol" else "List Bullet"
        for item in node.children:
            if item.tag == "li":
                _add_runs(document.add_paragraph(style=style), item)
        return

    if tag == "pre":
        paragraph = document.add_paragraph()
        run = paragraph.add_run("".join(child.text for child in _walk(node, "#text")))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        return

    if tag == "blockquote":
        for child in node.children:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            _add_runs(paragraph, child)
        return

    if tag == "hr":
        paragraph = document.add_paragraph()
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "c9d2dc")
            borders.append(bottom)
            p_pr.append(borders)
        except Exception:
            paragraph.add_run("—" * 20)
        return

    if tag == "table":
        _emit_table(document, node)
        return

    _add_runs(document.add_paragraph(), node)


def html_to_docx(html: str, path, *, title: str = "", base_dirs=()):
    """Export HTML (as produced by :func:`markdown_to_html`) to a Word document."""
    global _INLINE_BASE_DIRS

    from docx import Document

    document = Document()
    _INLINE_BASE_DIRS = tuple(base_dirs)
    _style_document(document)
    if title:
        heading = document.add_heading(title, level=0)
        heading.alignment = 1

    root = _parse_html(_extract_body(html))
    for element in root.children:
        _emit_block(document, element, base_dirs=tuple(base_dirs))

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def safe_stem(name: object, default: str = "paper") -> str:
    """Sanitise a caller-supplied filename stem so it cannot escape the output dir.

    ``stem``/``name`` come from a JSON spec that a model can author, so
    ``"../../etc/passwd"`` must not turn into a path traversal.
    """
    stem = re.sub(r"[^\w.\-]+", "_", str(name or "")).strip("._-") or default
    return stem[:80]


def _try_pdf(docx_path: Path) -> Path | None:
    """Best-effort docx -> pdf via a local LibreOffice, if present."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            "/usr/bin/soffice",
        ):
            if Path(candidate).is_file():
                soffice = candidate
                break
    if not soffice:
        return None
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir",
             str(docx_path.parent), str(docx_path)],
            timeout=180, check=True, capture_output=True,
        )
    except Exception:
        return None
    pdf_path = docx_path.with_suffix(".pdf")
    return pdf_path if pdf_path.is_file() else None


def build_paper(markdown_text: str, out_dir, *, title: str = "", stem: str = "paper",
                css: str = "", base_dirs=None, meta: dict[str, str] | None = None,
                pdf: bool = False) -> dict:
    """Write ``<out_dir>/<stem>.html`` and ``<stem>.docx`` (and ``.pdf`` if asked)."""
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    stem = safe_stem(stem)

    parsed_meta, body = parse_front_matter(markdown_text)
    merged = dict(parsed_meta)
    if meta:
        merged.update({key: str(value) for key, value in meta.items() if str(value).strip()})
    if title and not merged.get("title"):
        merged["title"] = title

    search_dirs = list(base_dirs or [])
    search_dirs.append(directory)
    search_dirs.append(Path.cwd())

    html = markdown_to_html(body, title=str(merged.get("title") or title), css=css, meta=merged)
    html_path = directory / f"{stem}.html"
    html_path.write_text(html, encoding="utf-8")
    # The title block is already part of the HTML, so do not add a second title.
    docx_path = html_to_docx(html, directory / f"{stem}.docx", base_dirs=search_dirs)

    result = {"html": html_path, "docx": docx_path}
    if pdf:
        generated = _try_pdf(docx_path)
        if generated is not None:
            result["pdf"] = generated
    return result
